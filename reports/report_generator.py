from io import BytesIO
from datetime import datetime

import pandas as pd

from docx import Document
from docx.shared import Inches, Pt

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def _safe(value, default="N/A"):
    """Return a safe display value."""
    if value is None:
        return default

    if isinstance(value, float):
        if pd.isna(value):
            return default

    return value


def _fmt(value, decimals=3):
    """Format numerical values safely."""
    if value is None:
        return "N/A"

    try:
        if pd.isna(value):
            return "N/A"
    except Exception:
        pass

    try:
        return f"{float(value):.{decimals}f}"
    except (TypeError, ValueError):
        return str(value)


def _get(reactor, *keys, default=None):
    """Return the first available key from a reactor dictionary."""
    for key in keys:
        if key in reactor and reactor[key] is not None:
            return reactor[key]

    return default


# ============================================================
# SUMMARY DATAFRAME
# ============================================================

def create_summary_dataframe(reactors):
    """
    Create a summary DataFrame for all reactors.

    Parameters
    ----------
    reactors : list[dict]
        Reactor calculation dictionaries.

    Returns
    -------
    pandas.DataFrame
    """

    rows = []

    for i, reactor in enumerate(reactors, start=1):

        row = {
            "Reactor": _get(
                reactor,
                "name",
                "reactor_name",
                default=f"Reactor {i}"
            ),

            "Volume (m³)": _get(
                reactor,
                "volume_m3",
                "working_volume_m3",
                default=None
            ),

            "Vessel Volume (m³)": _get(
                reactor,
                "vessel_volume_m3",
                "total_volume_m3",
                default=None
            ),

            "Tank Diameter (m)": _get(
                reactor,
                "tank_diameter_m",
                "diameter_m",
                "D",
                default=None
            ),

            "Liquid Height (m)": _get(
                reactor,
                "liquid_height_m",
                "HL",
                default=None
            ),

            "RPM": _get(
                reactor,
                "rpm",
                "speed_rpm",
                "N",
                default=None
            ),

            "Impeller Diameter (m)": _get(
                reactor,
                "impeller_diameter_m",
                "Di",
                default=None
            ),

            "Number of Impellers": _get(
                reactor,
                "number_impellers",
                "impellers",
                default=None
            ),

            "Power (kW)": _get(
                reactor,
                "power_kw",
                "shaft_power_kw",
                default=None
            ),

            "P/V (kW/m³)": _get(
                reactor,
                "power_volume_kw_m3",
                "P_V_kW_m3",
                default=None
            ),

            "P/V (W/m³)": _get(
                reactor,
                "power_volume_w_m3",
                "P_V_W_m3",
                default=None
            ),

            "Tip Speed (m/s)": _get(
                reactor,
                "tip_speed_m_s",
                "tip_speed",
                default=None
            ),

            "Reynolds Number": _get(
                reactor,
                "reynolds_number",
                "Re",
                default=None
            ),

            "Froude Number": _get(
                reactor,
                "froude_number",
                "Fr",
                default=None
            ),

            "Pumping Capacity (m³/h)": _get(
                reactor,
                "pumping_capacity_m3_h",
                "pumping_rate_m3_h",
                default=None
            ),

            "Turnover (1/h)": _get(
                reactor,
                "turnover_1_h",
                "turnover_rate_1_h",
                default=None
            ),

            "Njs (RPM)": _get(
                reactor,
                "njs_rpm",
                "Njs",
                default=None
            ),

            "kLa (1/h)": _get(
                reactor,
                "kLa_1_h",
                "kla_1_h",
                "kla",
                default=None
            ),

            "Gas Holdup": _get(
                reactor,
                "gas_holdup_fraction",
                "gas_holdup",
                default=None
            ),
        }

        rows.append(row)

    return pd.DataFrame(rows)


# ============================================================
# WORD REPORT
# ============================================================

def create_word_report(
    project_name="Reactor Scale-Up Study",
    prepared_by="",
    process_type="",
    study_mode="",
    scaleup_basis="",
    reactors=None,
    scaleup_result=None,
):
    """
    Generate a professional Word report.

    Returns
    -------
    BytesIO
        Word document stored in memory.
    """

    if reactors is None:
        reactors = []

    document = Document()

    # --------------------------------------------------------
    # DOCUMENT DEFAULT FONT
    # --------------------------------------------------------

    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9)

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    title = document.add_heading(
        "REACTOR SCALE-UP ENGINEERING REPORT",
        level=0
    )

    title.alignment = 1

    document.add_paragraph(
        f"Project: {_safe(project_name)}"
    )

    document.add_paragraph(
        f"Prepared By: {_safe(prepared_by)}"
    )

    document.add_paragraph(
        f"Process Type: {_safe(process_type)}"
    )

    document.add_paragraph(
        f"Study Mode: {_safe(study_mode)}"
    )

    document.add_paragraph(
        f"Scale-Up Basis: {_safe(scaleup_basis)}"
    )

    document.add_paragraph(
        f"Report Date: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    )

    document.add_paragraph("")

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    document.add_heading(
        "1. Executive Summary",
        level=1
    )

    document.add_paragraph(
        "This report presents the reactor scale-up engineering "
        "assessment including reactor geometry, agitation, "
        "mixing performance, power requirements, gas-liquid "
        "parameters where applicable, scale-up comparison and "
        "engineering validation."
    )

    # --------------------------------------------------------
    # REACTOR SUMMARY
    # --------------------------------------------------------

    document.add_heading(
        "2. Reactor Summary",
        level=1
    )

    summary_df = create_summary_dataframe(reactors)

    if not summary_df.empty:

        table = document.add_table(
            rows=1,
            cols=len(summary_df.columns)
        )

        table.style = "Table Grid"

        header_cells = table.rows[0].cells

        for j, column in enumerate(summary_df.columns):
            header_cells[j].text = str(column)

        for _, row in summary_df.iterrows():

            cells = table.add_row().cells

            for j, value in enumerate(row):
                cells[j].text = _fmt(value)

    else:

        document.add_paragraph(
            "No reactor calculation data available."
        )

    # --------------------------------------------------------
    # INDIVIDUAL REACTOR DETAILS
    # --------------------------------------------------------

    document.add_heading(
        "3. Reactor Design Details",
        level=1
    )

    for i, reactor in enumerate(reactors, start=1):

        reactor_name = _get(
            reactor,
            "name",
            "reactor_name",
            default=f"Reactor {i}"
        )

        document.add_heading(
            str(reactor_name),
            level=2
        )

        details = [
            (
                "Working Volume",
                _get(
                    reactor,
                    "volume_m3",
                    "working_volume_m3"
                ),
                "m³"
            ),
            (
                "Vessel Volume",
                _get(
                    reactor,
                    "vessel_volume_m3",
                    "total_volume_m3"
                ),
                "m³"
            ),
            (
                "Tank Diameter",
                _get(
                    reactor,
                    "tank_diameter_m",
                    "diameter_m",
                    "D"
                ),
                "m"
            ),
            (
                "Liquid Height",
                _get(
                    reactor,
                    "liquid_height_m",
                    "HL"
                ),
                "m"
            ),
            (
                "RPM",
                _get(
                    reactor,
                    "rpm",
                    "speed_rpm",
                    "N"
                ),
                "RPM"
            ),
            (
                "Impeller Diameter",
                _get(
                    reactor,
                    "impeller_diameter_m",
                    "Di"
                ),
                "m"
            ),
            (
                "Number of Impellers",
                _get(
                    reactor,
                    "number_impellers",
                    "impellers"
                ),
                "-"
            ),
            (
                "Power",
                _get(
                    reactor,
                    "power_kw",
                    "shaft_power_kw"
                ),
                "kW"
            ),
            (
                "P/V",
                _get(
                    reactor,
                    "power_volume_kw_m3",
                    "P_V_kW_m3"
                ),
                "kW/m³"
            ),
            (
                "P/V",
                _get(
                    reactor,
                    "power_volume_w_m3",
                    "P_V_W_m3"
                ),
                "W/m³"
            ),
            (
                "Tip Speed",
                _get(
                    reactor,
                    "tip_speed_m_s",
                    "tip_speed"
                ),
                "m/s"
            ),
            (
                "Reynolds Number",
                _get(
                    reactor,
                    "reynolds_number",
                    "Re"
                ),
                "-"
            ),
            (
                "Froude Number",
                _get(
                    reactor,
                    "froude_number",
                    "Fr"
                ),
                "-"
            ),
            (
                "Pumping Capacity",
                _get(
                    reactor,
                    "pumping_capacity_m3_h",
                    "pumping_rate_m3_h"
                ),
                "m³/h"
            ),
            (
                "Turnover",
                _get(
                    reactor,
                    "turnover_1_h",
                    "turnover_rate_1_h"
                ),
                "1/h"
            ),
            (
                "Njs",
                _get(
                    reactor,
                    "njs_rpm",
                    "Njs"
                ),
                "RPM"
            ),
            (
                "kLa",
                _get(
                    reactor,
                    "kLa_1_h",
                    "kla_1_h",
                    "kla"
                ),
                "1/h"
            ),
        ]

        table = document.add_table(
            rows=1,
            cols=3
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "Parameter"
        table.rows[0].cells[1].text = "Value"
        table.rows[0].cells[2].text = "Unit"

        for parameter, value, unit in details:

            cells = table.add_row().cells

            cells[0].text = str(parameter)
            cells[1].text = _fmt(value)
            cells[2].text = str(unit)

        document.add_paragraph("")

    # --------------------------------------------------------
    # SCALE-UP
    # --------------------------------------------------------

    if scaleup_result is not None:

        document.add_heading(
            "4. Scale-Up Assessment",
            level=1
        )

        if isinstance(scaleup_result, dict):

            table = document.add_table(
                rows=1,
                cols=3
            )

            table.style = "Table Grid"

            table.rows[0].cells[0].text = "Parameter"
            table.rows[0].cells[1].text = "Value"
            table.rows[0].cells[2].text = "Unit"

            for key, value in scaleup_result.items():

                cells = table.add_row().cells

                cells[0].text = str(key)

                if isinstance(value, (int, float)):
                    cells[1].text = _fmt(value)
                else:
                    cells[1].text = str(value)

                cells[2].text = "-"

        else:

            document.add_paragraph(
                str(scaleup_result)
            )

    # --------------------------------------------------------
    # ENGINEERING NOTES
    # --------------------------------------------------------

    document.add_heading(
        "5. Engineering Notes",
        level=1
    )

    document.add_paragraph(
        "The calculated results should be reviewed against "
        "process-specific operating requirements, laboratory "
        "or pilot-scale data, equipment vendor information and "
        "applicable engineering design standards before final "
        "equipment specification."
    )

    document.add_paragraph(
        "Particular attention should be given to agitator "
        "power, P/V, tip speed, suspension performance, "
        "gas dispersion, Njs, heat-transfer requirements and "
        "mechanical limitations during scale-up."
    )

    # --------------------------------------------------------
    # DISCLAIMER
    # --------------------------------------------------------

    document.add_heading(
        "6. Disclaimer",
        level=1
    )

    document.add_paragraph(
        "This report is intended for engineering study and "
        "preliminary design purposes. Final equipment design, "
        "mechanical design, agitator selection and process "
        "validation should be confirmed by qualified engineers "
        "and equipment vendors."
    )

    # --------------------------------------------------------
    # SAVE TO MEMORY
    # --------------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# ============================================================
# EXCEL REPORT
# ============================================================

def create_excel_report(
    project_name="Reactor Scale-Up Study",
    prepared_by="",
    process_type="",
    study_mode="",
    scaleup_basis="",
    reactors=None,
    scaleup_result=None,
):
    """
    Generate a professional Excel report.

    Returns
    -------
    BytesIO
        Excel workbook stored in memory.
    """

    if reactors is None:
        reactors = []

    workbook = Workbook()

    # --------------------------------------------------------
    # SUMMARY SHEET
    # --------------------------------------------------------

    ws_summary = workbook.active
    ws_summary.title = "Summary"

    ws_summary["A1"] = "REACTOR SCALE-UP ENGINEERING REPORT"
    ws_summary["A1"].font = Font(
        bold=True,
        size=16
    )

    ws_summary["A3"] = "Project"
    ws_summary["B3"] = _safe(project_name)

    ws_summary["A4"] = "Prepared By"
    ws_summary["B4"] = _safe(prepared_by)

    ws_summary["A5"] = "Process Type"
    ws_summary["B5"] = _safe(process_type)

    ws_summary["A6"] = "Study Mode"
    ws_summary["B6"] = _safe(study_mode)

    ws_summary["A7"] = "Scale-Up Basis"
    ws_summary["B7"] = _safe(scaleup_basis)

    ws_summary["A8"] = "Report Date"
    ws_summary["B8"] = datetime.now().strftime(
        "%d-%m-%Y %H:%M"
    )

    summary_df = create_summary_dataframe(reactors)

    if not summary_df.empty:

        start_row = 11

        for col_idx, column in enumerate(
            summary_df.columns,
            start=1
        ):

            cell = ws_summary.cell(
                row=start_row,
                column=col_idx
            )

            cell.value = column
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_idx, row in enumerate(
            summary_df.itertuples(index=False),
            start=start_row + 1
        ):

            for col_idx, value in enumerate(
                row,
                start=1
            ):

                ws_summary.cell(
                    row=row_idx,
                    column=col_idx,
                    value=_safe(value, "")
                )

    # --------------------------------------------------------
    # INDIVIDUAL REACTOR SHEETS
    # --------------------------------------------------------

    for i, reactor in enumerate(reactors, start=1):

        reactor_name = _get(
            reactor,
            "name",
            "reactor_name",
            default=f"Reactor {i}"
        )

        # Excel sheet names have restrictions.
        sheet_name = f"Reactor {i}"

        ws = workbook.create_sheet(
            title=sheet_name
        )

        ws["A1"] = str(reactor_name)

        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        headers = [
            "Parameter",
            "Value",
            "Unit"
        ]

        for col_idx, header in enumerate(
            headers,
            start=1
        ):

            cell = ws.cell(
                row=3,
                column=col_idx
            )

            cell.value = header
            cell.font = Font(bold=True)

        details = [
            (
                "Working Volume",
                _get(
                    reactor,
                    "volume_m3",
                    "working_volume_m3"
                ),
                "m³"
            ),
            (
                "Vessel Volume",
                _get(
                    reactor,
                    "vessel_volume_m3",
                    "total_volume_m3"
                ),
                "m³"
            ),
            (
                "Tank Diameter",
                _get(
                    reactor,
                    "tank_diameter_m",
                    "diameter_m",
                    "D"
                ),
                "m"
            ),
            (
                "Liquid Height",
                _get(
                    reactor,
                    "liquid_height_m",
                    "HL"
                ),
                "m"
            ),
            (
                "RPM",
                _get(
                    reactor,
                    "rpm",
                    "speed_rpm",
                    "N"
                ),
                "RPM"
            ),
            (
                "Impeller Diameter",
                _get(
                    reactor,
                    "impeller_diameter_m",
                    "Di"
                ),
                "m"
            ),
            (
                "Number of Impellers",
                _get(
                    reactor,
                    "number_impellers",
                    "impellers"
                ),
                "-"
            ),
            (
                "Power",
                _get(
                    reactor,
                    "power_kw",
                    "shaft_power_kw"
                ),
                "kW"
            ),
            (
                "P/V",
                _get(
                    reactor,
                    "power_volume_kw_m3",
                    "P_V_kW_m3"
                ),
                "kW/m³"
            ),
            (
                "P/V",
                _get(
                    reactor,
                    "power_volume_w_m3",
                    "P_V_W_m3"
                ),
                "W/m³"
            ),
            (
                "Tip Speed",
                _get(
                    reactor,
                    "tip_speed_m_s",
                    "tip_speed"
                ),
                "m/s"
            ),
            (
                "Reynolds Number",
                _get(
                    reactor,
                    "reynolds_number",
                    "Re"
                ),
                "-"
            ),
            (
                "Froude Number",
                _get(
                    reactor,
                    "froude_number",
                    "Fr"
                ),
                "-"
            ),
            (
                "Pumping Capacity",
                _get(
                    reactor,
                    "pumping_capacity_m3_h",
                    "pumping_rate_m3_h"
                ),
                "m³/h"
            ),
            (
                "Turnover",
                _get(
                    reactor,
                    "turnover_1_h",
                    "turnover_rate_1_h"
                ),
                "1/h"
            ),
            (
                "Njs",
                _get(
                    reactor,
                    "njs_rpm",
                    "Njs"
                ),
                "RPM"
            ),
            (
                "kLa",
                _get(
                    reactor,
                    "kLa_1_h",
                    "kla_1_h",
                    "kla"
                ),
                "1/h"
            ),
        ]

        for row_idx, (
            parameter,
            value,
            unit
        ) in enumerate(
            details,
            start=4
        ):

            ws.cell(
                row=row_idx,
                column=1,
                value=parameter
            )

            ws.cell(
                row=row_idx,
                column=2,
                value=_safe(value, "")
            )

            ws.cell(
                row=row_idx,
                column=3,
                value=unit
            )

    # --------------------------------------------------------
    # SCALE-UP SHEET
    # --------------------------------------------------------

    if scaleup_result is not None:

        ws_scale = workbook.create_sheet(
            title="Scale-Up"
        )

        ws_scale["A1"] = "SCALE-UP ASSESSMENT"

        ws_scale["A1"].font = Font(
            bold=True,
            size=14
        )

        ws_scale["A3"] = "Parameter"
        ws_scale["B3"] = "Value"
        ws_scale["C3"] = "Unit"

        for cell in ws_scale[3]:

            cell.font = Font(
                bold=True
            )

        if isinstance(scaleup_result, dict):

            row = 4

            for key, value in scaleup_result.items():

                ws_scale.cell(
                    row=row,
                    column=1,
                    value=str(key)
                )

                ws_scale.cell(
                    row=row,
                    column=2,
                    value=_safe(value, "")
                )

                ws_scale.cell(
                    row=row,
                    column=3,
                    value="-"
                )

                row += 1

        else:

            ws_scale["A4"] = str(
                scaleup_result
            )

    # --------------------------------------------------------
    # ENGINEERING NOTES
    # --------------------------------------------------------

    ws_notes = workbook.create_sheet(
        title="Engineering Notes"
    )

    ws_notes["A1"] = "ENGINEERING NOTES"

    ws_notes["A1"].font = Font(
        bold=True,
        size=14
    )

    notes = [
        "Review reactor geometry against actual equipment drawings.",
        "Confirm agitator selection with the agitator vendor.",
        "Check P/V and tip speed against process requirements.",
        "Check solids suspension and Njs where applicable.",
        "For gas-liquid systems, validate gas dispersion and kLa.",
        "Confirm heat-transfer requirements independently.",
        "Final mechanical design must be completed separately.",
        "Use pilot or plant data wherever available for scale-up validation.",
    ]

    for i, note in enumerate(
        notes,
        start=3
    ):

        ws_notes.cell(
            row=i,
            column=1,
            value=note
        )

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    for ws in workbook.worksheets:

        for column_cells in ws.columns:

            try:

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                max_length = 0

                for cell in column_cells:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    40
                )

            except Exception:
                pass

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output
