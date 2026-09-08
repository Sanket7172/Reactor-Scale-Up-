from io import BytesIO
from datetime import datetime
import json
import math

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

from docx import Document


# ============================================================
# VALUE HELPERS
# ============================================================

def _excel_value(value):

    if value is None:
        return ""

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        (int, float),
    ):

        if isinstance(
            value,
            float,
        ) and not math.isfinite(
            value
        ):

            return ""

        return value

    if isinstance(
        value,
        (dict, list, tuple),
    ):

        return json.dumps(
            value,
            default=str,
            ensure_ascii=False,
        )

    return str(value)


def _flatten_dict(
    data,
    prefix="",
):

    rows = []

    if not isinstance(
        data,
        dict,
    ):

        return rows

    for key, value in data.items():

        name = (
            f"{prefix}.{key}"
            if prefix
            else str(key)
        )

        if isinstance(
            value,
            dict,
        ):

            rows.extend(
                _flatten_dict(
                    value,
                    name,
                )
            )

        elif isinstance(
            value,
            (list, tuple),
        ):

            rows.append(
                (
                    name,
                    _excel_value(value),
                )
            )

        else:

            rows.append(
                (
                    name,
                    value,
                )
            )

    return rows


# ============================================================
# EXCEL FORMAT
# ============================================================

def _format_sheet(
    worksheet,
):

    thin = Side(
        style="thin",
        color="D9D9D9",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            cell.border = border

    for column in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            try:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    ),
                )

            except Exception:
                pass

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12,
            ),
            50,
        )

    worksheet.freeze_panes = "A4"


def _add_title(
    worksheet,
    title,
):

    worksheet["A1"] = title

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A1"].alignment = Alignment(
        vertical="center"
    )

    worksheet["A2"] = (
        "Generated: "
        + datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )


def _add_headers(
    worksheet,
    row,
    headers,
):

    fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for index, header in enumerate(
        headers,
        1,
    ):

        cell = worksheet.cell(
            row=row,
            column=index,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = fill


# ============================================================
# EXCEL REPORT
# ============================================================

def create_excel_report(
    inputs,
    result,
):

    if not isinstance(
        inputs,
        dict,
    ):

        raise ValueError(
            "inputs must be a dictionary."
        )

    if not isinstance(
        result,
        dict,
    ):

        raise ValueError(
            "result must be a dictionary."
        )

    output = BytesIO()

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    ws = workbook.create_sheet(
        "Summary"
    )

    _add_title(
        ws,
        "REACTOR SCALE-UP & MIXING CALCULATION REPORT",
    )

    _add_headers(
        ws,
        4,
        [
            "Parameter",
            "Value",
            "Unit",
        ],
    )

    summary = [
        (
            "Working Volume",
            result.get("volume_m3"),
            "m³",
        ),
        (
            "Vessel Volume",
            result.get("vessel_volume_m3"),
            "m³",
        ),
        (
            "Tank Diameter",
            result.get("tank_diameter_m"),
            "m",
        ),
        (
            "Straight Height",
            result.get("straight_height_m"),
            "m",
        ),
        (
            "Liquid Height",
            result.get("liquid_height_m"),
            "m",
        ),
        (
            "RPM",
            result.get("rpm"),
            "rpm",
        ),
        (
            "Impeller",
            result.get("agitator"),
            "-",
        ),
        (
            "Impeller Diameter",
            result.get("impeller_diameter_m"),
            "m",
        ),
        (
            "Number of Impellers",
            result.get("number_impellers"),
            "-",
        ),
        (
            "Reynolds Number",
            result.get("reynolds_number"),
            "-",
        ),
        (
            "Froude Number",
            result.get("froude_number"),
            "-",
        ),
        (
            "Tip Speed",
            result.get("tip_speed"),
            "m/s",
        ),
        (
            "Power",
            result.get("power_kw"),
            "kW",
        ),
        (
            "Specific Power",
            result.get(
                "power_volume_kw_m3"
            ),
            "kW/m³",
        ),
        (
            "Torque",
            result.get("torque_nm"),
            "N·m",
        ),
        (
            "Pumping Capacity",
            result.get("pumping_m3_h"),
            "m³/h",
        ),
        (
            "Turnover Time",
            result.get(
                "turnover_time_min"
            ),
            "min",
        ),
        (
            "Mixing Regime",
            result.get(
                "mixing_regime"
            ),
            "-",
        ),
    ]

    row = 5

    for parameter, value, unit in summary:

        ws.cell(
            row=row,
            column=1,
            value=parameter,
        )

        ws.cell(
            row=row,
            column=2,
            value=_excel_value(value),
        )

        ws.cell(
            row=row,
            column=3,
            value=unit,
        )

        row += 1

    # ========================================================
    # INPUTS
    # ========================================================

    ws = workbook.create_sheet(
        "Inputs"
    )

    _add_title(
        ws,
        "REACTOR INPUTS",
    )

    _add_headers(
        ws,
        4,
        [
            "Parameter",
            "Value",
        ],
    )

    row = 5

    for key, value in _flatten_dict(
        inputs
    ):

        # Skip detailed impeller list here.
        if key == "impellers":
            continue

        ws.cell(
            row=row,
            column=1,
            value=key,
        )

        ws.cell(
            row=row,
            column=2,
            value=_excel_value(value),
        )

        row += 1

    # ========================================================
    # IMPELLERS
    # ========================================================

    ws = workbook.create_sheet(
        "Impellers"
    )

    _add_title(
        ws,
        "IMPELLER ARRANGEMENT",
    )

    _add_headers(
        ws,
        4,
        [
            "No.",
            "Position",
            "Agitator",
            "Flow Type",
            "Diameter (m)",
            "D/T",
            "Elevation (m)",
            "Bottom Clearance (m)",
            "Np",
            "Nq",
        ],
    )

    impellers = inputs.get(
        "impellers",
        []
    )

    tank_D = float(
        result.get(
            "tank_diameter_m",
            0.0,
        )
        or 0.0
    )

    row = 5

    for index, impeller in enumerate(
        impellers,
        1,
    ):

        if not isinstance(
            impeller,
            dict,
        ):
            continue

        agitator = impeller.get(
            "agitator_type",
            impeller.get(
                "type",
                "",
            ),
        )

        diameter = float(
            impeller.get(
                "diameter_m",
                impeller.get(
                    "D",
                    0.0,
                ),
            )
            or 0.0
        )

        D_T = (
            diameter / tank_D
            if tank_D > 0
            else None
        )

        try:

            from libraries.agitator_geometry import (
                AGITATORS,
            )

            data = AGITATORS.get(
                agitator,
                {},
            )

            flow_type = data.get(
                "flow_type",
                "",
            )

            np_value = data.get(
                "np"
            )

            nq_value = data.get(
                "nq"
            )

        except Exception:

            flow_type = ""
            np_value = None
            nq_value = None

        values = [
            index,
            impeller.get(
                "position",
                "",
            ),
            agitator,
            flow_type,
            diameter,
            D_T,
            impeller.get(
                "elevation_m"
            ),
            impeller.get(
                "bottom_clearance_m"
            ),
            np_value,
            nq_value,
        ]

        for column, value in enumerate(
            values,
            1,
        ):

            ws.cell(
                row=row,
                column=column,
                value=_excel_value(value),
            )

        row += 1

    # ========================================================
    # PERFORMANCE
    # ========================================================

    ws = workbook.create_sheet(
        "Performance"
    )

    _add_title(
        ws,
        "MIXING PERFORMANCE",
    )

    _add_headers(
        ws,
        4,
        [
            "Parameter",
            "Value",
            "Unit",
        ],
    )

    performance = [
        (
            "Reynolds Number",
            result.get(
                "reynolds_number"
            ),
            "-",
        ),
        (
            "Froude Number",
            result.get(
                "froude_number"
            ),
            "-",
        ),
        (
            "Tip Speed",
            result.get(
                "tip_speed"
            ),
            "m/s",
        ),
        (
            "Power",
            result.get(
                "power_kw"
            ),
            "kW",
        ),
        (
            "Specific Power",
            result.get(
                "power_volume_kw_m3"
            ),
            "kW/m³",
        ),
        (
            "Torque",
            result.get(
                "torque_nm"
            ),
            "N·m",
        ),
        (
            "Pumping Capacity",
            result.get(
                "pumping_m3_h"
            ),
            "m³/h",
        ),
        (
            "Pumping / Volume",
            result.get(
                "pumping_per_volume"
            ),
            "1/h",
        ),
        (
            "Turnover Time",
            result.get(
                "turnover_time_min"
            ),
            "min",
        ),
        (
            "D/T",
            result.get(
                "D_T"
            ),
            "-",
        ),
        (
            "H/T",
            result.get(
                "H_T"
            ),
            "-",
        ),
        (
            "Mixing Regime",
            result.get(
                "mixing_regime"
            ),
            "-",
        ),
    ]

    row = 5

    for parameter, value, unit in performance:

        ws.cell(
            row=row,
            column=1,
            value=parameter,
        )

        ws.cell(
            row=row,
            column=2,
            value=_excel_value(value),
        )

        ws.cell(
            row=row,
            column=3,
            value=unit,
        )

        row += 1

    # ========================================================
    # GAS-LIQUID
    # ========================================================

    ws = workbook.create_sheet(
        "Gas-Liquid"
    )

    _add_title(
        ws,
        "GAS-LIQUID MASS TRANSFER",
    )

    _add_headers(
        ws,
        4,
        [
            "Parameter",
            "Value",
            "Unit",
        ],
    )

    gas_rows = [
        (
            "Gas Flow",
            result.get(
                "gas_flow_m3_h"
            ),
            "m³/h",
        ),
        (
            "Superficial Velocity",
            result.get(
                "gas_superficial_velocity_m_s"
            ),
            "m/s",
        ),
        (
            "Gas Holdup",
            result.get(
                "gas_holdup_fraction"
            ),
            "-",
        ),
        (
            "Bubble Diameter",
            result.get(
                "bubble_diameter_m"
            ),
            "m",
        ),
        (
            "Bubble Rise Velocity",
            result.get(
                "bubble_rise_velocity_m_s"
            ),
            "m/s",
        ),
        (
            "Bubble Residence Time",
            result.get(
                "bubble_residence_time_s"
            ),
            "s",
        ),
        (
            "Bubble Reynolds",
            result.get(
                "bubble_reynolds"
            ),
            "-",
        ),
        (
            "Schmidt Number",
            result.get(
                "schmidt_number"
            ),
            "-",
        ),
        (
            "Sherwood Number",
            result.get(
                "sherwood_number"
            ),
            "-",
        ),
        (
            "kL",
            result.get(
                "kL_m_s"
            ),
            "m/s",
        ),
        (
            "Interfacial Area",
            result.get(
                "interfacial_area_m2_m3"
            ),
            "m²/m³",
        ),
        (
            "kLa",
            result.get(
                "kLa_1_s"
            ),
            "1/s",
        ),
        (
            "kLa",
            result.get(
                "kLa_1_h"
            ),
            "1/h",
        ),
        (
            "Status",
            result.get(
                "gas_liquid_status"
            ),
            "-",
        ),
    ]

    row = 5

    for parameter, value, unit in gas_rows:

        ws.cell(
            row=row,
            column=1,
            value=parameter,
        )

        ws.cell(
            row=row,
            column=2,
            value=_excel_value(value),
        )

        ws.cell(
            row=row,
            column=3,
            value=unit,
        )

        row += 1

    # ========================================================
    # CALCULATION DATA
    # ========================================================

    ws = workbook.create_sheet(
        "Calculation Data"
    )

    _add_title(
        ws,
        "ENGINE CALCULATION OUTPUT",
    )

    _add_headers(
        ws,
        4,
        [
            "Parameter",
            "Value",
        ],
    )

    row = 5

    for key, value in _flatten_dict(
        result
    ):

        ws.cell(
            row=row,
            column=1,
            value=key,
        )

        ws.cell(
            row=row,
            column=2,
            value=_excel_value(value),
        )

        row += 1

    # ========================================================
    # FORMAT ALL SHEETS
    # ========================================================

    for worksheet in workbook.worksheets:

        _format_sheet(
            worksheet
        )

    workbook.save(
        output
    )

    output.seek(0)

    return output


# ============================================================
# WORD REPORT
# ============================================================

def create_word_report(
    inputs,
    result,
):

    if not isinstance(
        inputs,
        dict,
    ):

        raise ValueError(
            "inputs must be a dictionary."
        )

    if not isinstance(
        result,
        dict,
    ):

        raise ValueError(
            "result must be a dictionary."
        )

    output = BytesIO()

    document = Document()

    document.add_heading(
        "Reactor Scale-Up & Mixing Calculation Report",
        level=0,
    )

    document.add_paragraph(
        "Generated: "
        + datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )

    # ========================================================
    # EXECUTIVE SUMMARY
    # ========================================================

    document.add_heading(
        "1. Executive Summary",
        level=1,
    )

    summary_table = document.add_table(
        rows=1,
        cols=3,
    )

    summary_table.style = "Table Grid"

    headers = [
        "Parameter",
        "Value",
        "Unit",
    ]

    for i, header in enumerate(
        headers
    ):

        summary_table.rows[
            0
        ].cells[i].text = header

    summary = [
        (
            "Working Volume",
            result.get("volume_m3"),
            "m³",
        ),
        (
            "Tank Diameter",
            result.get(
                "tank_diameter_m"
            ),
            "m",
        ),
        (
            "Liquid Height",
            result.get(
                "liquid_height_m"
            ),
            "m",
        ),
        (
            "RPM",
            result.get(
                "rpm"
            ),
            "rpm",
        ),
        (
            "Agitator",
            result.get(
                "agitator"
            ),
            "-",
        ),
        (
            "Impeller Diameter",
            result.get(
                "impeller_diameter_m"
            ),
            "m",
        ),
        (
            "Power",
            result.get(
                "power_kw"
            ),
            "kW",
        ),
        (
            "Specific Power",
            result.get(
                "power_volume_kw_m3"
            ),
            "kW/m³",
        ),
        (
            "Reynolds Number",
            result.get(
                "reynolds_number"
            ),
            "-",
        ),
        (
            "Mixing Regime",
            result.get(
                "mixing_regime"
            ),
            "-",
        ),
    ]

    for parameter, value, unit in summary:

        cells = (
            summary_table.add_row().cells
        )

        cells[0].text = str(
            parameter
        )

        cells[1].text = str(
            _excel_value(value)
        )

        cells[2].text = unit

    # ========================================================
    # INPUTS
    # ========================================================

    document.add_heading(
        "2. Reactor Inputs",
        level=1,
    )

    input_table = document.add_table(
        rows=1,
        cols=2,
    )

    input_table.style = "Table Grid"

    input_table.rows[
        0
    ].cells[0].text = "Parameter"

    input_table.rows[
        0
    ].cells[1].text = "Value"

    for key, value in _flatten_dict(
        inputs
    ):

        if key == "impellers":
            continue

        cells = (
            input_table.add_row().cells
        )

        cells[0].text = str(
            key
        )

        cells[1].text = str(
            _excel_value(value)
        )

    # ========================================================
    # IMPELLERS
    # ========================================================

    document.add_heading(
        "3. Impeller Arrangement",
        level=1,
    )

    impellers = inputs.get(
        "impellers",
        []
    )

    if impellers:

        table = document.add_table(
            rows=1,
            cols=7,
        )

        table.style = "Table Grid"

        headers = [
            "No.",
            "Position",
            "Agitator",
            "Diameter (m)",
            "D/T",
            "Elevation (m)",
            "Clearance (m)",
        ]

        for i, header in enumerate(
            headers
        ):

            table.rows[
                0
            ].cells[i].text = header

        tank_D = float(
            result.get(
                "tank_diameter_m",
                0.0,
            )
            or 0.0
        )

        for index, imp in enumerate(
            impellers,
            1,
        ):

            if not isinstance(
                imp,
                dict,
            ):
                continue

            D = float(
                imp.get(
                    "diameter_m",
                    imp.get(
                        "D",
                        0.0,
                    ),
                )
                or 0.0
            )

            D_T = (
                D / tank_D
                if tank_D > 0
                else 0.0
            )

            values = [
                index,
                imp.get(
                    "position",
                    "",
                ),
                imp.get(
                    "agitator_type",
                    imp.get(
                        "type",
                        "",
                    ),
                ),
                f"{D:.3f}",
                f"{D_T:.3f}",
                f"{float(imp.get('elevation_m', 0.0) or 0.0):.3f}",
                f"{float(imp.get('bottom_clearance_m', 0.0) or 0.0):.3f}",
            ]

            cells = (
                table.add_row().cells
            )

            for i, value in enumerate(
                values
            ):

                cells[i].text = str(
                    value
                )

    # ========================================================
    # PERFORMANCE
    # ========================================================

    document.add_heading(
        "4. Mixing Performance",
        level=1,
    )

    performance_table = document.add_table(
        rows=1,
        cols=3,
    )

    performance_table.style = "Table Grid"

    for i, header in enumerate(
        headers[:3]
    ):

        performance_table.rows[
            0
        ].cells[i].text = header

    rows = [
        (
            "Reynolds Number",
            result.get(
                "reynolds_number"
            ),
            "-",
        ),
        (
            "Froude Number",
            result.get(
                "froude_number"
            ),
            "-",
        ),
        (
            "Tip Speed",
            result.get(
                "tip_speed"
            ),
            "m/s",
        ),
        (
            "Power",
            result.get(
                "power_kw"
            ),
            "kW",
        ),
        (
            "Specific Power",
            result.get(
                "power_volume_kw_m3"
            ),
            "kW/m³",
        ),
        (
            "Torque",
            result.get(
                "torque_nm"
            ),
            "N·m",
        ),
        (
            "Pumping Capacity",
            result.get(
                "pumping_m3_h"
            ),
            "m³/h",
        ),
        (
            "Turnover Time",
            result.get(
                "turnover_time_min"
            ),
            "min",
        ),
        (
            "Mixing Regime",
            result.get(
                "mixing_regime"
            ),
            "-",
        ),
    ]

    for parameter, value, unit in rows:

        cells = (
            performance_table.add_row().cells
        )

        cells[0].text = parameter
        cells[1].text = str(
            _excel_value(value)
        )
        cells[2].text = unit

    # ========================================================
    # GAS-LIQUID
    # ========================================================

    document.add_heading(
        "5. Gas–Liquid Mass Transfer",
        level=1,
    )

    gas_rows = [
        (
            "Gas Flow",
            result.get(
                "gas_flow_m3_h"
            ),
            "m³/h",
        ),
        (
            "Gas Holdup",
            result.get(
                "gas_holdup_fraction"
            ),
            "-",
        ),
        (
            "Bubble Diameter",
            result.get(
                "bubble_diameter_m"
            ),
            "m",
        ),
        (
            "Bubble Reynolds",
            result.get(
                "bubble_reynolds"
            ),
            "-",
        ),
        (
            "Schmidt Number",
            result.get(
                "schmidt_number"
            ),
            "-",
        ),
        (
            "Sherwood Number",
            result.get(
                "sherwood_number"
            ),
            "-",
        ),
        (
            "kL",
            result.get(
                "kL_m_s"
            ),
            "m/s",
        ),
        (
            "Interfacial Area",
            result.get(
                "interfacial_area_m2_m3"
            ),
            "m²/m³",
        ),
        (
            "kLa",
            result.get(
                "kLa_1_h"
            ),
            "1/h",
        ),
    ]

    table = document.add_table(
        rows=1,
        cols=3,
    )

    table.style = "Table Grid"

    for i, header in enumerate(
        headers[:3]
    ):

        table.rows[
            0
        ].cells[i].text = header

    for parameter, value, unit in gas_rows:

        cells = (
            table.add_row().cells
        )

        cells[0].text = parameter
        cells[1].text = str(
            _excel_value(value)
        )
        cells[2].text = unit

    # ========================================================
    # ENGINEERING NOTE
    # ========================================================

    document.add_heading(
        "6. Engineering Basis & Limitations",
        level=1,
    )

    document.add_paragraph(
        "The calculations in this report are intended "
        "for preliminary process engineering screening, "
        "reactor scale-up assessment and equipment "
        "comparison. Agitator power correlations, "
        "gas-liquid mass-transfer correlations and "
        "geometry calculations should be validated against "
        "vendor data, plant data, pilot trials or "
        "appropriate published correlations before final "
        "equipment design."
    )

    document.add_paragraph(
        "The current multi-impeller calculation engine "
        "uses the selected primary impeller correlation "
        "with the configured number of impellers. "
        "Different impeller-specific power correlations "
        "in the same shaft require an enhanced "
        "multi-impeller model."
    )

    document.save(
        output
    )

    output.seek(0)

    return output
